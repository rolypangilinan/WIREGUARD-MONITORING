"""Generate MATERIAL_SETTER_SEQUENCE.xlsx from the sequence data."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = openpyxl.Workbook()

# Styles
title_font = Font(name="Calibri", size=14, bold=True)
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
section_font = Font(name="Calibri", size=12, bold=True)
section_fill = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
note_font = Font(name="Calibri", size=10, italic=True, color="FF0000")
cell_font = Font(name="Calibri", size=10)
thin_border = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin")
)
wrap_align = Alignment(wrap_text=True, vertical="center")

def add_header_row(ws, row, headers):
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border

def add_data_row(ws, row, values):
    for col, v in enumerate(values, 1):
        c = ws.cell(row=row, column=col, value=v)
        c.font = cell_font
        c.alignment = wrap_align
        c.border = thin_border

def add_section_title(ws, row, title, col_span=4):
    c = ws.cell(row=row, column=1, value=title)
    c.font = section_font
    c.fill = section_fill
    c.border = thin_border
    for col in range(2, col_span + 1):
        cc = ws.cell(row=row, column=col)
        cc.fill = section_fill
        cc.border = thin_border

def add_note(ws, row, text, col_span=4):
    c = ws.cell(row=row, column=1, value=text)
    c.font = note_font
    c.alignment = wrap_align

# ============ SHEET 1: ALL SEQUENCES ============
ws = wb.active
ws.title = "Sequences"
ws.column_dimensions["A"].width = 8
ws.column_dimensions["B"].width = 65
ws.column_dimensions["C"].width = 16
ws.column_dimensions["D"].width = 35

r = 1
ws.cell(row=r, column=1, value="MATERIAL SETTER — FULL SEQUENCE").font = title_font
r += 1
ws.cell(row=r, column=1, value="Type 1 (Pin-Controlled): Arduino pin directly controls relay. Active LOW = ON. Pins: 4, 6, 10, 12.").font = Font(name="Calibri", size=10, italic=True)
r += 1
ws.cell(row=r, column=1, value="Type 2 (PC-Commanded): PC sends command to Arduino, Arduino sets pin. Used for interlocks. Pins: 3, 8, 13, 27. HIGH = open.").font = Font(name="Calibri", size=10, italic=True)
r += 2

# --- SEQUENCE A ---
add_section_title(ws, r, "SEQUENCE A: Operator Places ID at Material Setter (MS ID Scanner)")
r += 1
add_header_row(ws, r, ["Step", "Action", "Pin", "Direction"])
r += 1
seq_a = [
    [1, "IR Sensor 1 detects ID (LOW)", "Pin 2", "SENSOR → ARDUINO"],
    [2, "Arduino sends MS_LOAD_ID", "—", "ARDUINO → PC"],
    [3, "Arduino turns ON Scanner 1 relay (LOW)", "Pin 4", "ARDUINO PIN (Type 1)"],
    [4, "Arduino waits 2s boot delay", "—", "INTERNAL"],
    [5, "Arduino sends trigger {0x16,0x54,0x0D,0x0A} to R35D-B via Serial1", "Pin 18 TX", "ARDUINO → SCANNER"],
    [6, "Arduino sends MS_SCANNER_READY", "—", "ARDUINO → PC"],
    [7, "R35D-B scans QR, sends data back via Serial1", "Pin 19 RX", "SCANNER → ARDUINO"],
    [8, "Arduino sends MS_SCAN:<data>", "—", "ARDUINO → PC"],
    [9, "Arduino sends stop cmd to scanner", "Pin 18 TX", "ARDUINO → SCANNER"],
    [10, "Arduino turns OFF Scanner 1 relay (HIGH)", "Pin 4", "ARDUINO PIN (Type 1)"],
]
for row_data in seq_a:
    add_data_row(ws, r, row_data)
    r += 1
r += 1
add_note(ws, r, "IF ID IS REMOVED:")
r += 1
seq_a_remove = [
    [11, "IR Sensor 1 goes HIGH (200ms sustained)", "Pin 2", "SENSOR → ARDUINO"],
    [12, "Arduino cancels scan, stops scanner, relay OFF", "Pin 4", "ARDUINO PIN (Type 1)"],
    [13, "Arduino sends MS_UNLOAD_ID", "—", "ARDUINO → PC"],
]
for row_data in seq_a_remove:
    add_data_row(ws, r, row_data)
    r += 1
r += 2

# --- SEQUENCE B ---
add_section_title(ws, r, "SEQUENCE B: Kitting Dunpla Placed at Material Setter (MS Kitting Scanner)")
r += 1
add_header_row(ws, r, ["Step", "Action", "Pin", "Direction"])
r += 1
seq_b = [
    [1, "IR Sensor 5 detects dunpla (LOW)", "Pin 5", "SENSOR → ARDUINO"],
    [2, "Arduino sends MS_KITTING_LOAD", "—", "ARDUINO → PC"],
    [3, "Arduino turns ON Kitting Scanner relay (LOW)", "Pin 6", "ARDUINO PIN (Type 1)"],
    [4, "Arduino waits 2s boot delay", "—", "INTERNAL"],
    [5, "Arduino sends trigger to Kitting R35D-B via Serial2", "Pin 16 TX", "ARDUINO → SCANNER"],
    [6, "Arduino sends MS_KIT_SCANNER_READY", "—", "ARDUINO → PC"],
    [7, "R35D-B scans QR, sends data back via Serial2", "Pin 17 RX", "SCANNER → ARDUINO"],
    [8, "Arduino sends MS_KITSCAN:<data>", "—", "ARDUINO → PC"],
    [9, "Arduino sends stop cmd to scanner", "Pin 16 TX", "ARDUINO → SCANNER"],
    [10, "Arduino turns OFF Kitting Scanner relay (HIGH)", "Pin 6", "ARDUINO PIN (Type 1)"],
]
for row_data in seq_b:
    add_data_row(ws, r, row_data)
    r += 1
r += 1
add_note(ws, r, "PC PROCESSES KITTING DATA:")
r += 1
seq_b_pc = [
    [11, "PC (Flask) receives MS_KITSCAN:<data>, saves to DB", "—", "PC INTERNAL"],
    [12, "PC queues RELAY2_ON command", "—", "PC → ARDUINO (queued)"],
    [13, "Arduino receives RELAY2_ON, sets Pin 3 HIGH", "Pin 3", "PC → ARDUINO → RELAY (Type 2)"],
    [14, "MS solenoid interlock OPENS (dunpla can pass)", "Pin 3", "SOLENOID OPENS"],
    [15, "Arduino sends MS_OPEN_INTERLOCK", "—", "ARDUINO → PC"],
]
for row_data in seq_b_pc:
    add_data_row(ws, r, row_data)
    r += 1
r += 1
add_note(ws, r, "DUNPLA REMOVED FROM KITTING SENSOR:")
r += 1
seq_b_remove = [
    [16, "IR Sensor 5 goes HIGH (500ms sustained)", "Pin 5", "SENSOR → ARDUINO"],
    [17, "Arduino sends MS_KITTING_UNLOAD", "—", "ARDUINO → PC"],
    [18, "Arduino powers OFF kitting scanner relay", "Pin 6", "ARDUINO PIN (Type 1)"],
]
for row_data in seq_b_remove:
    add_data_row(ws, r, row_data)
    r += 1
r += 2

# --- SEQUENCE C ---
add_section_title(ws, r, "SEQUENCE C: Dunpla Reaches Standby Position")
r += 1
add_header_row(ws, r, ["Step", "Action", "Pin", "Direction"])
r += 1
seq_c = [
    [1, "IR Sensor 2 detects dunpla (LOW, 2s sustained)", "Pin 7", "SENSOR → ARDUINO"],
    [2, "Arduino sends STANDBY_MS", "—", "ARDUINO → PC"],
    [3, "PC queues RELAY2_OFF (close MS interlock)", "—", "PC → ARDUINO (queued)"],
    [4, "Arduino receives RELAY2_OFF, sets Pin 3 LOW", "Pin 3", "PC → ARDUINO → RELAY (Type 2)"],
    [5, "MS solenoid interlock CLOSES", "Pin 3", "SOLENOID CLOSES"],
    [6, "Arduino sends MS_CLOSE_INTERLOCK", "—", "ARDUINO → PC"],
    [7, "PC queues STANDBY_RELAY_ON", "—", "PC → ARDUINO (queued)"],
    [8, "Arduino receives cmd, sets Pin 8 HIGH", "Pin 8", "PC → ARDUINO → RELAY (Type 2)"],
    [9, "Standby interlock OPENS (dunpla can go to P1)", "Pin 8", "SOLENOID OPENS"],
    [10, "Arduino sends STANDBY_OPEN_INTERLOCK", "—", "ARDUINO → PC"],
]
for row_data in seq_c:
    add_data_row(ws, r, row_data)
    r += 1
r += 2

# --- SEQUENCE D ---
add_section_title(ws, r, "SEQUENCE D: Dunpla Reaches Process 1")
r += 1
add_header_row(ws, r, ["Step", "Action", "Pin", "Direction"])
r += 1
seq_d = [
    [1, "IR Sensor 6 detects dunpla (LOW, 500ms sustained)", "Pin 22", "SENSOR → ARDUINO"],
    [2, "Arduino sends P1_DONE", "—", "ARDUINO → PC"],
    [3, "PC queues STANDBY_RELAY_OFF", "—", "PC → ARDUINO (queued)"],
    [4, "Arduino receives cmd, sets Pin 8 LOW", "Pin 8", "PC → ARDUINO → RELAY (Type 2)"],
    [5, "Standby interlock CLOSES", "Pin 8", "SOLENOID CLOSES"],
    [6, "Arduino sends STANDBY_CLOSE_INTERLOCK", "—", "ARDUINO → PC"],
]
for row_data in seq_d:
    add_data_row(ws, r, row_data)
    r += 1
r += 2

# --- SEQUENCE E ---
add_section_title(ws, r, "SEQUENCE E: Operator Places ID at Process 1 (P1 ID Scanner 3)")
r += 1
add_header_row(ws, r, ["Step", "Action", "Pin", "Direction"])
r += 1
seq_e = [
    [1, "IR Sensor 3 detects ID (LOW, 500ms sustained)", "Pin 9", "SENSOR → ARDUINO"],
    [2, "Arduino sends P1_ID_LOAD", "—", "ARDUINO → PC"],
    [3, "Arduino turns ON P1 Scanner 3 relay (LOW)", "Pin 10", "ARDUINO PIN (Type 1)"],
    [4, "Arduino waits 2s boot delay", "—", "INTERNAL"],
    [5, "Arduino sends trigger to P1 R35D-B via Serial3", "Pin 14 TX", "ARDUINO → SCANNER"],
    [6, "Arduino sends P1_SCANNER_READY", "—", "ARDUINO → PC"],
    [7, "R35D-B scans QR, sends data back via Serial3", "Pin 15 RX", "SCANNER → ARDUINO"],
    [8, "Arduino sends P1_SCAN:<data>", "—", "ARDUINO → PC"],
    [9, "Arduino sends stop cmd to scanner", "Pin 14 TX", "ARDUINO → SCANNER"],
    [10, "Arduino turns OFF P1 Scanner 3 relay (HIGH)", "Pin 10", "ARDUINO PIN (Type 1)"],
]
for row_data in seq_e:
    add_data_row(ws, r, row_data)
    r += 1
r += 1
add_note(ws, r, "IF ID IS REMOVED:")
r += 1
seq_e_remove = [
    [11, "IR Sensor 3 goes HIGH (500ms sustained)", "Pin 9", "SENSOR → ARDUINO"],
    [12, "Arduino cancels scan, stops scanner, relay OFF", "Pin 10", "ARDUINO PIN (Type 1)"],
    [13, "Arduino sends P1_ID_UNLOAD", "—", "ARDUINO → PC"],
]
for row_data in seq_e_remove:
    add_data_row(ws, r, row_data)
    r += 1
r += 2

# --- SEQUENCE F ---
add_section_title(ws, r, "SEQUENCE F: P1 Kitting Dunpla Placed (P1 Kitting Scanner 4)")
r += 1
add_note(ws, r, "NOTE: Sensor & relay on MS Arduino. Scanner serial data on CT Arduino (Serial1).")
r += 1
add_header_row(ws, r, ["Step", "Action", "Pin", "Direction"])
r += 1
seq_f = [
    [1, "IR Sensor detects dunpla (LOW)", "Pin 11", "SENSOR → MS ARDUINO"],
    [2, "MS Arduino sends P1_KITTING_LOAD", "—", "MS ARDUINO → PC"],
    [3, "MS Arduino turns ON Scanner 4 relay (LOW). Scan cycle LOCKED.", "Pin 12", "MS ARDUINO PIN (Type 1)"],
    [4, "MS Arduino waits 2s boot delay", "—", "INTERNAL"],
    [5, "MS Arduino sends P1_KIT_SCANNER_READY", "—", "MS ARDUINO → PC"],
    [6, "PC queues P1_KIT_SCAN_TRIGGER for CT Arduino", "—", "PC → CT ARDUINO (queued)"],
    [7, "CT Arduino sends trigger to Scanner 4 via Serial1", "Pin 18 TX", "CT ARDUINO → SCANNER"],
    [8, "R35D-B scans QR, sends data back via Serial1", "Pin 19 RX", "SCANNER → CT ARDUINO"],
    [9, "CT Arduino sends P1_KITSCAN:<data>", "—", "CT ARDUINO → PC"],
    [10, "PC receives data, queues P1_KIT_SCAN_DONE for MS Arduino", "—", "PC → MS ARDUINO (queued)"],
    [11, "MS Arduino receives P1_KIT_SCAN_DONE, ends scan cycle", "—", "PC → MS ARDUINO"],
    [12, "MS Arduino turns OFF Scanner 4 relay (HIGH)", "Pin 12", "MS ARDUINO PIN (Type 1)"],
    [13, "MS Arduino sends P1_KIT_SCAN_DONE_CONFIRM", "—", "MS ARDUINO → PC"],
]
for row_data in seq_f:
    add_data_row(ws, r, row_data)
    r += 1
r += 1
add_note(ws, r, "SAFETY TIMEOUT (if scan never completes, 10s):")
r += 1
seq_f_timeout = [
    ["T", "MS Arduino auto-powers OFF relay after 10s", "Pin 12", "MS ARDUINO PIN (Type 1)"],
    ["T+1", "MS Arduino sends P1_KIT_SCAN_TIMEOUT", "—", "MS ARDUINO → PC"],
]
for row_data in seq_f_timeout:
    add_data_row(ws, r, row_data)
    r += 1
r += 2

# --- SEQUENCE G ---
add_section_title(ws, r, "SEQUENCE G: P2 ID Detection")
r += 1
add_header_row(ws, r, ["Step", "Action", "Pin", "Direction"])
r += 1
seq_g = [
    [1, "IR Sensor 7 detects ID (LOW, 500ms sustained)", "Pin 23", "SENSOR → ARDUINO"],
    [2, "Arduino sends P2_ID_LOAD", "—", "ARDUINO → PC"],
    [3, "(removal) IR Sensor 7 goes HIGH (500ms)", "Pin 23", "SENSOR → ARDUINO"],
    [4, "Arduino sends P2_ID_UNLOAD", "—", "ARDUINO → PC"],
]
for row_data in seq_g:
    add_data_row(ws, r, row_data)
    r += 1
r += 2

# --- SEQUENCE H ---
add_section_title(ws, r, "SEQUENCE H: Dunpla Reaches Process 2")
r += 1
add_header_row(ws, r, ["Step", "Action", "Pin", "Direction"])
r += 1
seq_h = [
    [1, "IR Sensor 8 detects dunpla (LOW, 500ms sustained)", "Pin 28", "SENSOR → ARDUINO"],
    [2, "Arduino sends P2_DONE", "—", "ARDUINO → PC"],
]
for row_data in seq_h:
    add_data_row(ws, r, row_data)
    r += 1
r += 2

# --- SEQUENCE I ---
add_section_title(ws, r, "SEQUENCE I: P1 Interlock (PC-Commanded)")
r += 1
add_header_row(ws, r, ["Action", "Pin", "Direction", ""])
r += 1
seq_i = [
    ["PC queues RELAY_P1_ON", "—", "PC → ARDUINO", ""],
    ["Arduino sets Pin 13 HIGH", "Pin 13", "PC → ARDUINO → RELAY (Type 2)", ""],
    ["Arduino sends P1_OPEN_INTERLOCK", "—", "ARDUINO → PC", ""],
    ["PC queues RELAY_P1_OFF", "—", "PC → ARDUINO", ""],
    ["Arduino sets Pin 13 LOW", "Pin 13", "PC → ARDUINO → RELAY (Type 2)", ""],
    ["Arduino sends P1_CLOSE_INTERLOCK", "—", "ARDUINO → PC", ""],
]
for row_data in seq_i:
    add_data_row(ws, r, row_data)
    r += 1
r += 2

# --- SEQUENCE J ---
add_section_title(ws, r, "SEQUENCE J: P2 Interlock (PC-Commanded)")
r += 1
add_header_row(ws, r, ["Action", "Pin", "Direction", ""])
r += 1
seq_j = [
    ["PC queues RELAY_P2_ON", "—", "PC → ARDUINO", ""],
    ["Arduino sets Pin 27 HIGH", "Pin 27", "PC → ARDUINO → RELAY (Type 2)", ""],
    ["Arduino sends P2_OPEN_INTERLOCK", "—", "ARDUINO → PC", ""],
    ["PC queues RELAY_P2_OFF", "—", "PC → ARDUINO", ""],
    ["Arduino sets Pin 27 LOW", "Pin 27", "PC → ARDUINO → RELAY (Type 2)", ""],
    ["Arduino sends P2_CLOSE_INTERLOCK", "—", "ARDUINO → PC", ""],
]
for row_data in seq_j:
    add_data_row(ws, r, row_data)
    r += 1
r += 2

# ============ SHEET 2: PIN SUMMARY ============
ws2 = wb.create_sheet("Pin Summary")
ws2.column_dimensions["A"].width = 8
ws2.column_dimensions["B"].width = 14
ws2.column_dimensions["C"].width = 16
ws2.column_dimensions["D"].width = 45
ws2.column_dimensions["E"].width = 40
ws2.column_dimensions["F"].width = 60

ws2.cell(row=1, column=1, value="PIN SUMMARY TABLE").font = title_font
add_header_row(ws2, 3, ["Pin", "INPUT / OUTPUT", "INPUT_PULLUP", "Function", "Relay Type", "Purpose & Usage"])

pins = [
    ["2", "INPUT", "YES", "IR Sensor 1 — MS ID detection", "—",
     "Detects operator ID card at Material Setter. LOW=ID present. Triggers MS_LOAD_ID and powers ON Scanner 1 (Pin 4) to scan the ID QR code."],
    ["3", "OUTPUT", "NO", "Relay 2 — MS solenoid interlock", "Type 2 (PC-Commanded, HIGH=open)",
     "Controls MS gate solenoid. PC sends RELAY2_ON after kitting QR is saved to open gate. PC sends RELAY2_OFF when dunpla reaches Standby (STANDBY_MS). Initial state: LOW (closed)."],
    ["4", "OUTPUT", "NO", "Relay 4 — MS ID scanner power", "Type 1 (Pin-Controlled, LOW=on)",
     "Powers the MS ID R35D-B scanner. Arduino sets LOW when ID detected (Pin 2), sets HIGH after scan complete or ID removed. Initial state: HIGH (OFF)."],
    ["5", "INPUT", "YES", "IR Sensor 5 — MS kitting detection", "—",
     "Detects kitting dunpla at Material Setter. LOW=dunpla present. Triggers MS_KITTING_LOAD and powers ON Kitting Scanner (Pin 6) to scan kitting QR code."],
    ["6", "OUTPUT", "NO", "Relay 5 — MS kitting scanner power", "Type 1 (Pin-Controlled, LOW=on)",
     "Powers the MS Kitting R35D-B scanner. Arduino sets LOW when dunpla detected (Pin 5), sets HIGH after scan complete or dunpla removed. Initial state: HIGH (OFF)."],
    ["7", "INPUT", "YES", "IR Sensor 2 — dunpla at Standby", "—",
     "Detects dunpla arrival at Standby position. LOW sustained 2s = STANDBY_MS. Triggers PC to close MS interlock (RELAY2_OFF on Pin 3) and open Standby interlock (STANDBY_RELAY_ON on Pin 8)."],
    ["8", "OUTPUT", "NO", "Standby Relay — Standby interlock", "Type 2 (PC-Commanded, HIGH=open)",
     "Controls Standby gate solenoid. PC sends STANDBY_RELAY_ON when dunpla reaches Standby. PC sends STANDBY_RELAY_OFF when dunpla reaches Process 1 (P1_DONE). Initial state: LOW (closed)."],
    ["9", "INPUT", "YES", "IR Sensor 3 — P1 ID detection", "—",
     "Detects operator ID card at Process 1. LOW sustained 500ms = P1_ID_LOAD. Triggers power ON of P1 Scanner 3 (Pin 10) to scan operator ID. HIGH sustained 500ms = P1_ID_UNLOAD."],
    ["10", "OUTPUT", "NO", "Relay 7 — P1 Scanner 3 power", "Type 1 (Pin-Controlled, LOW=on)",
     "Powers the P1 ID R35D-B scanner (Scanner 3). Arduino sets LOW when ID detected (Pin 9), sets HIGH after scan complete or ID removed. Initial state: HIGH (OFF)."],
    ["11", "INPUT", "YES", "IR Sensor — P1 kitting detection", "—",
     "Detects kitting dunpla at Process 1. LOW = P1_KITTING_LOAD. Triggers power ON of P1 Kitting Scanner 4 relay (Pin 12). Scanner serial data is read by CT Arduino (not this board)."],
    ["12", "OUTPUT", "NO", "Relay 8 — P1 Scanner 4 power", "Type 1 (Pin-Controlled, LOW=on)",
     "Powers the P1 Kitting R35D-B scanner (Scanner 4). Relay is LOCKED during scan cycle to survive sensor glitches. Unlocked by P1_KIT_SCAN_DONE from PC or 10s safety timeout. Initial state: HIGH (OFF)."],
    ["13", "OUTPUT", "NO", "Relay 9 — P1 solenoid interlock", "Type 2 (PC-Commanded, HIGH=open)",
     "Controls Process 1 gate solenoid. PC sends RELAY_P1_ON to open, RELAY_P1_OFF to close. Initial state: LOW (closed)."],
    ["14", "TX3", "NO", "P1 Scanner 3 trigger (Serial3 TX)", "—",
     "Arduino sends trigger command {0x16,0x54,0x0D,0x0A} to P1 ID R35D-B scanner via Serial3 TX. Also sends stop command {0x16,0x55,0x0D,0x0A} after successful scan."],
    ["15", "RX3", "NO", "P1 Scanner 3 data (Serial3 RX)", "—",
     "Receives scanned QR data from P1 ID R35D-B scanner via Serial3 RX at 115200 baud. Data forwarded as P1_SCAN:<data> to PC."],
    ["16", "TX2", "NO", "MS Kitting Scanner trigger (Serial2 TX)", "—",
     "Arduino sends trigger command to MS Kitting R35D-B scanner via Serial2 TX. Also sends stop command after successful scan."],
    ["17", "RX2", "NO", "MS Kitting Scanner data (Serial2 RX)", "—",
     "Receives scanned QR data from MS Kitting R35D-B scanner via Serial2 RX at 115200 baud. Data forwarded as MS_KITSCAN:<data> to PC."],
    ["18", "TX1", "NO", "MS ID Scanner trigger (Serial1 TX)", "—",
     "Arduino sends trigger command to MS ID R35D-B scanner via Serial1 TX. Also sends stop command after successful scan."],
    ["19", "RX1", "NO", "MS ID Scanner data (Serial1 RX)", "—",
     "Receives scanned QR data from MS ID R35D-B scanner via Serial1 RX at 115200 baud. Data forwarded as MS_SCAN:<data> to PC."],
    ["22", "INPUT", "YES", "IR Sensor 6 — dunpla at Process 1", "—",
     "Detects dunpla arrival at Process 1. LOW sustained 500ms = P1_DONE. Triggers PC to close Standby interlock (STANDBY_RELAY_OFF on Pin 8) and stop P1 server timer."],
    ["23", "INPUT", "YES", "IR Sensor 7 — P2 ID detection", "—",
     "Detects operator ID card at Process 2. LOW sustained 500ms = P2_ID_LOAD. HIGH sustained 500ms = P2_ID_UNLOAD. No scanner attached — ID detection only."],
    ["27", "OUTPUT", "NO", "Relay — P2 solenoid interlock", "Type 2 (PC-Commanded, HIGH=open)",
     "Controls Process 2 gate solenoid. PC sends RELAY_P2_ON to open, RELAY_P2_OFF to close. Initial state: LOW (closed)."],
    ["28", "INPUT", "YES", "IR Sensor 8 — dunpla at Process 2", "—",
     "Detects dunpla arrival at Process 2. LOW sustained 500ms = P2_DONE. Triggers PC to stop P2 server timer."],
]
for i, pin_data in enumerate(pins):
    add_data_row(ws2, 4 + i, pin_data)

# Save
output_path = r"c:\Users\Admin\Desktop\NEW HP 60_80 LINE\CYCLE TIME MONITORING\SAVE\NEW\4 W_ SUFFIX AND QTY  BUT W_INFO AT PY TERMINAL\3 HAVE 2 KINDS OF DATABASE\CHAT GPT\3 60-80HP-LINE-APP 4A - FIXED POLLING END POINTS\MATERIAL_SETTER_SEQUENCE.xlsx"
wb.save(output_path)
print(f"Excel saved to: {output_path}")
